"""Bounded local document operations. Input/output are JSON; sources are never overwritten."""
import datetime as dt
import hashlib
import json
import os
from pathlib import Path
import re
import shutil
import sys
import tempfile


def fingerprint(path):
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def file_path(value, existing=True):
    path = Path(str(value or ""))
    if not path.is_absolute():
        raise ValueError("Cần đường dẫn tuyệt đối tới file user đã cung cấp.")
    path = path.resolve()
    if existing and (not path.is_file() or path.stat().st_size > 100 * 1024 * 1024):
        raise ValueError("File không tồn tại hoặc lớn hơn 100 MB.")
    if path.suffix.lower() not in {".docx", ".xlsx", ".xlsm", ".pdf"}:
        raise ValueError("Hỗ trợ .docx, .xlsx, .xlsm và .pdf; file .doc/.xls cần chuyển định dạng.")
    return path


def serial(value):
    if isinstance(value, (dt.datetime, dt.date, dt.time)):
        return value.isoformat()
    return value


def word_blocks(document):
    seen = set()

    def container_blocks(container, prefix):
        for index, paragraph in enumerate(container.paragraphs):
            if paragraph._p not in seen:
                seen.add(paragraph._p)
                yield f"{prefix}/p{index}", paragraph
        for index, table in enumerate(container.tables):
            for row_index, row in enumerate(table.rows):
                for column_index, cell in enumerate(row.cells):
                    yield from container_blocks(cell, f"{prefix}/t{index}/r{row_index}/c{column_index}")

    yield from container_blocks(document, "body")
    for index, section in enumerate(document.sections):
        for name in ("header", "footer", "first_page_header", "first_page_footer",
                     "even_page_header", "even_page_footer"):
            container = getattr(section, name)
            # Accessing an undefined header/footer can create it in python-docx.
            if container._has_definition:
                yield from container_blocks(container, f"section{index}/{name}")


def open_workbook(path, read_only=False, data_only=False):
    from openpyxl import load_workbook
    return load_workbook(path, read_only=read_only, data_only=data_only,
                         keep_vba=path.suffix.lower() == ".xlsm", keep_links=True)


def inspect(path):
    info = {"path": str(path), "format": path.suffix.lower()[1:],
            "bytes": path.stat().st_size, "fingerprint": fingerprint(path)}
    if path.suffix.lower() == ".docx":
        from docx import Document
        document = Document(path)
        info.update(block_count=sum(1 for _ in word_blocks(document)), table_count=len(document.tables))
    elif path.suffix.lower() in {".xlsx", ".xlsm"}:
        workbook = open_workbook(path, read_only=True)
        try:
            info["sheets"] = [{"name": s.title, "rows": s.max_row, "columns": s.max_column}
                              for s in workbook.worksheets]
        finally:
            workbook.close()
    else:
        from pypdf import PdfReader
        with path.open("rb") as stream:
            document = PdfReader(stream)
            if document.is_encrypted and not document.decrypt(""):
                raise ValueError("PDF cần mật khẩu; chưa đọc được nội dung.")
            info["page_count"] = len(document.pages)
    return info


def read_document(path, request):
    result = inspect(path)
    limit = max(1000, min(60000, int(request.get("max_chars", 20000))))
    start = max(0, int(request.get("start", 0)))
    text_offset = max(0, int(request.get("text_offset", 0)))
    count = max(1, min(100, int(request.get("count", 30))))
    blocks, used, truncated = [], 0, False

    def add(block):
        nonlocal used, truncated
        size = len(json.dumps(block, ensure_ascii=False, default=serial))
        if used + size > limit:
            truncated = True
            return False
        blocks.append(block)
        used += size
        return True

    suffix = path.suffix.lower()
    if suffix == ".docx":
        from docx import Document
        source = list(word_blocks(Document(path)))
        if request.get("block_id"):
            matches = [(index, block) for index, block in enumerate(source) if block[0] == request["block_id"]]
            if not matches:
                raise ValueError("Không tìm thấy block_id; hãy đọc lại cấu trúc file.")
            start = matches[0][0]
            count = 1
        for block_id, paragraph in source[start:start + count]:
            text = paragraph.text[text_offset:] if not blocks else paragraph.text
            if not add({"id": block_id, "text": text}):
                # Return a bounded preview of a long paragraph without losing its ID.
                if not blocks:
                    blocks.append({"id": block_id, "text": text[:limit // 2], "text_truncated": True,
                                   "next_text_offset": text_offset + limit // 2})
                break
        result["next_start"] = start + len(blocks) if start + len(blocks) < len(source) else None
    elif suffix in {".xlsx", ".xlsm"}:
        from openpyxl.utils.cell import range_boundaries
        workbook = open_workbook(path, read_only=True)
        cached = open_workbook(path, read_only=True, data_only=True)
        try:
            sheet = workbook[str(request.get("sheet") or workbook.sheetnames[0])]
            cell_range = request.get("range") or f"A{start + 1}:J{start + min(count, 50)}"
            left, top, right, bottom = range_boundaries(cell_range)
            if None in (left, top, right, bottom) or min(left, top) < 1 or right < left or bottom < top:
                raise ValueError("Cần vùng ô hữu hạn, ví dụ A1:F30.")
            if right > 16384 or bottom > 1048576 or (right - left + 1) * (bottom - top + 1) > 2000:
                raise ValueError("Mỗi lần đọc tối đa 2000 ô.")
            result.update(sheet=sheet.title, range=cell_range, values_recalculated=False)
            cached_rows = cached[sheet.title].iter_rows(min_row=top, max_row=bottom, min_col=left, max_col=right)
            for row, cached_row in zip(sheet.iter_rows(min_row=top, max_row=bottom, min_col=left, max_col=right), cached_rows):
                for column, (cell, cached_cell) in enumerate(zip(row, cached_row), left):
                    from openpyxl.utils import get_column_letter
                    address = f"{get_column_letter(column)}{cell.row}" if cell.value is not None else None
                    if cell.value is None:
                        continue
                    block = {"cell": address, "value": serial(cell.value)}
                    if cell.data_type == "f":
                        block["cached_value"] = serial(cached_cell.value)
                    if getattr(cell, "hyperlink", None):
                        block["hyperlink"] = cell.hyperlink.target
                    if not add(block):
                        break
                if truncated:
                    break
        finally:
            workbook.close()
            cached.close()
    else:
        from pypdf import PdfReader
        with path.open("rb") as stream:
            document = PdfReader(stream)
            if document.is_encrypted:
                document.decrypt("")
            page_count = len(document.pages)
            page_numbers = request.get("pages") or list(range(start + 1, min(start + 3, page_count) + 1))
            if len(page_numbers) > 20 or any(type(p) is not int or not 1 <= p <= page_count for p in page_numbers):
                raise ValueError("Chọn tối đa 20 trang hợp lệ, đánh số từ 1.")
            empty = []
            for page_number in page_numbers:
                full_text = document.pages[page_number - 1].extract_text() or ""
                if not full_text.strip():
                    empty.append(page_number)
                text = full_text[text_offset:]
                if not add({"page": page_number, "text": text}):
                    if not blocks:
                        blocks.append({"page": page_number, "text": text[:limit // 2], "text_truncated": True,
                                       "next_text_offset": text_offset + limit // 2})
                    break
            result.update(empty_pages=empty, needs_ocr=bool(empty))
    result.update(blocks=blocks, truncated=truncated)
    return result


def replace_word(paragraph, before, after):
    runs = paragraph.runs
    text = "".join(run.text for run in runs)
    if paragraph.text != text:
        raise ValueError("Đoạn có hyperlink/field phức tạp; cần công cụ Word chuyên dụng để giữ cấu trúc.")
    positions = [m.start() for m in re.finditer(re.escape(before), text)]
    for start in reversed(positions):
        end, offset, first, last = start + len(before), 0, None, None
        for index, run in enumerate(runs):
            next_offset = offset + len(run.text)
            if first is None and offset <= start < next_offset:
                first = (index, start - offset)
            if offset < end <= next_offset:
                last = (index, end - offset)
                break
            offset = next_offset
        if first is None or last is None:
            raise ValueError("Không xác định được vị trí thay thế trong đoạn.")
        i, begin = first
        j, finish = last
        prefix, suffix = runs[i].text[:begin], runs[j].text[finish:]
        if i == j:
            runs[i].text = prefix + after + suffix
        else:
            runs[i].text = prefix + after
            for index in range(i + 1, j):
                runs[index].text = ""
            runs[j].text = suffix


def save_result(output, save, source=None, expected=None):
    if not output.parent.is_dir() or output.exists():
        raise ValueError("Thư mục output phải tồn tại và tên file output chưa được dùng.")
    handle, temporary = tempfile.mkstemp(prefix=".simi-document-", suffix=output.suffix, dir=output.parent)
    os.close(handle)
    temporary = Path(temporary)
    try:
        save(temporary)
        info = inspect(temporary)  # Reopen and validate the generated document before publishing it.
        if source and fingerprint(source) != expected:
            raise ValueError("File nguồn đã thay đổi trong lúc sửa; hãy đọc lại rồi thực hiện lại.")
        created = False
        try:
            with output.open("xb") as destination:
                created = True
                with temporary.open("rb") as source_stream:
                    shutil.copyfileobj(source_stream, destination)
                destination.flush()
                os.fsync(destination.fileno())
        except Exception:
            if created:
                output.unlink(missing_ok=True)
            raise
        info["path"] = str(output)
        return info
    finally:
        temporary.unlink(missing_ok=True)


def edit_document(path, request):
    expected = request.get("fingerprint")
    if not expected or fingerprint(path) != expected:
        raise ValueError("File đã thay đổi hoặc chưa đọc. Đọc file và dùng fingerprint mới nhất.")
    output = file_path(request.get("output_path"), existing=False)
    if output == path or output.suffix.lower() != path.suffix.lower():
        raise ValueError("Lưu bản mới cùng định dạng; không ghi đè file nguồn.")
    edits = request.get("edits") or []
    if not 1 <= len(edits) <= 200:
        raise ValueError("Cần từ 1 tới 200 thay đổi cụ thể.")
    changes = 0
    suffix = path.suffix.lower()
    if suffix == ".docx":
        from docx import Document
        document = Document(path)
        blocks = dict(word_blocks(document))
        for edit in edits:
            block = blocks.get(edit.get("block_id"))
            before, after = edit.get("before"), edit.get("after")
            if block is None or not isinstance(before, str) or not before or not isinstance(after, str):
                raise ValueError("Thay đổi Word cần block_id, before và after chính xác.")
            count = block.text.count(before)
            if count != edit.get("expected_count", 1):
                raise ValueError("Số vị trí Word khớp không đúng; hãy đọc lại đoạn trước khi sửa.")
            replace_word(block, before, after)
            changes += count
        result = save_result(output, document.save, path, expected)
    elif suffix in {".xlsx", ".xlsm"}:
        workbook = open_workbook(path)
        try:
            for edit in edits:
                sheet = workbook[edit["sheet"]]
                address = str(edit.get("cell") or "")
                if not re.fullmatch(r"[A-Za-z]{1,3}[1-9][0-9]{0,6}", address):
                    raise ValueError("Mỗi thay đổi Excel cần một địa chỉ ô chính xác.")
                cell = sheet[address]
                if "expected_value" not in edit or serial(cell.value) != edit["expected_value"]:
                    raise ValueError(f"Giá trị hiện tại của {sheet.title}!{address} không khớp.")
                if "value" not in edit or not isinstance(edit["value"], (str, int, float, bool, type(None))):
                    raise ValueError("Giá trị ô mới phải là text, số, boolean hoặc null.")
                cell.value = edit["value"]
                changes += 1
            result = save_result(output, workbook.save, path, expected)
            result["values_recalculated"] = False
        finally:
            workbook.close()
    else:
        from pypdf import PdfReader, PdfWriter
        with path.open("rb") as stream:
            reader = PdfReader(stream)
            if reader.is_encrypted and not reader.decrypt(""):
                raise ValueError("PDF cần mật khẩu; chưa sửa được nội dung.")
            document = PdfWriter(clone_from=reader)
            for edit in edits:
                if edit.get("operation") == "select_pages":
                    pages = edit.get("pages") or []
                    if not pages or len(pages) > 10000 or any(type(p) is not int or not 1 <= p <= len(document.pages) for p in pages):
                        raise ValueError("Danh sách trang PDF không hợp lệ.")
                    selected = PdfWriter()
                    for page in pages:
                        selected.add_page(document.pages[page - 1])
                    document = selected
                elif edit.get("operation") == "rotate_page":
                    page, degrees = edit.get("page"), edit.get("degrees")
                    if type(page) is not int or not 1 <= page <= len(document.pages) or degrees not in (0, 90, 180, 270):
                        raise ValueError("Trang/góc xoay PDF không hợp lệ.")
                    document.pages[page - 1].rotation = degrees
                else:
                    raise ValueError("PDF hỗ trợ chọn/sắp xếp trang và xoay trang. Sửa chữ cần file Word nguồn hoặc công cụ PDF chuyên dụng.")
                changes += 1
            result = save_result(output, document.write, path, expected)
    result.update(changes=changes, source_path=str(path), source_preserved=True)
    return result


def main():
    request = json.load(sys.stdin)
    path = file_path(request.get("path"))
    action = request.get("action")
    if action == "inspect":
        return inspect(path)
    if action == "read":
        return read_document(path, request)
    if action == "edit":
        return edit_document(path, request)
    raise ValueError("Thao tác tài liệu không hợp lệ.")


if __name__ == "__main__":
    for stream in (sys.stdin, sys.stdout, sys.stderr):
        if hasattr(stream, "reconfigure"):
            stream.reconfigure(encoding="utf-8")
    try:
        print(json.dumps(main(), ensure_ascii=False, default=serial, allow_nan=False))
    except Exception as error:
        print(json.dumps({"error": {"message": str(error)}}, ensure_ascii=False))
        sys.exit(1)
