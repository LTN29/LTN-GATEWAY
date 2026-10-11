"""Real document roundtrips in a temporary directory, invoked by node:test."""
import hashlib
import json
from pathlib import Path
import subprocess
import sys
from docx import Document
from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill
from pypdf import PdfReader, PdfWriter

root, script = Path(sys.argv[1]), sys.argv[2]
sys.stdout.reconfigure(encoding="utf-8")


def call(action, path, fail=False, **options):
    result = subprocess.run([sys.executable, script], input=json.dumps(
        {"action": action, "path": str(path), **options}), encoding="utf-8", capture_output=True)
    payload = json.loads(result.stdout)
    assert (result.returncode != 0) == fail, result.stdout + result.stderr
    return payload


word = root / "Tiếng Việt nguồn.docx"
document = Document()
paragraph = document.add_paragraph()
paragraph.add_run("Chào ").bold = True
paragraph.add_run("anh chị ").italic = True
paragraph.add_run("em, chào anh chị em.")
document.add_table(rows=1, cols=1).cell(0, 0).text = "Nội dung bảng"
document.sections[0].header.paragraphs[0].text = "Tiêu đề công ty"
document.save(word)
source_bytes = word.read_bytes()
read = call("read", word)
assert read["blocks"][0]["text"] == "Chào anh chị em, chào anh chị em."
assert any(block["text"] == "Nội dung bảng" for block in read["blocks"])
assert any(block["text"] == "Tiêu đề công ty" for block in read["blocks"])
edited = root / "Kết quả Word.docx"
options = {"fingerprint": read["fingerprint"], "output_path": str(edited), "edits": [
    {"block_id": read["blocks"][0]["id"], "before": "anh chị em", "after": "mọi người", "expected_count": 2}]}
call("edit", word, **options)
updated = Document(edited)
assert updated.paragraphs[0].text == "Chào mọi người, chào mọi người."
assert updated.paragraphs[0].runs[0].bold
assert updated.paragraphs[0].runs[1].italic
assert word.read_bytes() == source_bytes
assert "error" in call("edit", word, fail=True, **options)  # Existing output is not overwritten.
call("edit", word, fail=True, **{**options, "output_path": str(root / "stale.docx"), "fingerprint": "bad"})
call("edit", word, fail=True, **{**options, "output_path": str(word)})
assert not (root / "stale.docx").exists()

long_word = root / "Đoạn dài.docx"
long_document = Document()
long_document.add_paragraph("a" * 1500 + "Phần cuối")
long_document.save(long_word)
first_chunk = call("read", long_word, max_chars=1000)
block = first_chunk["blocks"][0]
assert first_chunk["truncated"] is True and block["text_truncated"] is True
last_chunk = call("read", long_word, block_id=block["id"], text_offset=1500, max_chars=1000)
assert last_chunk["blocks"][0]["text"] == "Phần cuối"

excel = root / "Dữ liệu.xlsx"
workbook = Workbook()
sheet = workbook.active
sheet.title = "Báo cáo"
sheet["A1"] = "Doanh thu"
sheet["B1"] = 100
sheet["C1"] = "=B1*2"
sheet["B1"].fill = PatternFill("solid", fgColor="FF0000")
workbook.save(excel)
source_bytes = excel.read_bytes()
read = call("read", excel, sheet="Báo cáo", range="A1:C2")
assert next(b for b in read["blocks"] if b["cell"] == "C1")["value"] == "=B1*2"
assert read["values_recalculated"] is False
excel_out = root / "Dữ liệu mới.xlsx"
call("edit", excel, fingerprint=read["fingerprint"], output_path=str(excel_out), edits=[
    {"sheet": "Báo cáo", "cell": "B1", "expected_value": 100, "value": 200}])
workbook = load_workbook(excel_out)
assert workbook.active["B1"].value == 200
assert workbook.active["B1"].fill.fgColor.rgb == "00FF0000"
assert workbook.active["C1"].value == "=B1*2"
workbook.close()
assert excel.read_bytes() == source_bytes
call("read", excel, fail=True, range="A1:XFD100")
bad_out = root / "Sai.xlsx"
call("edit", excel, fail=True, fingerprint=read["fingerprint"], output_path=str(bad_out), edits=[
    {"sheet": "Báo cáo", "cell": "B1", "expected_value": 999, "value": 200}])
assert not bad_out.exists()

pdf = root / "Scan.pdf"
writer = PdfWriter()
writer.add_blank_page(width=200, height=300)
writer.add_blank_page(width=400, height=500)
writer.write(pdf)
source_bytes = pdf.read_bytes()
read = call("read", pdf)
assert read["needs_ocr"] is True and read["empty_pages"] == [1, 2]
pdf_out = root / "Trang chọn.pdf"
call("edit", pdf, fingerprint=read["fingerprint"], output_path=str(pdf_out), edits=[
    {"operation": "select_pages", "pages": [2]}, {"operation": "rotate_page", "page": 1, "degrees": 90}])
reader = PdfReader(pdf_out)
assert len(reader.pages) == 1
assert reader.pages[0].rotation == 90
assert reader.pages[0].mediabox.width == 400
assert pdf.read_bytes() == source_bytes

from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
text_pdf = root / "Có text.pdf"
writer = PdfWriter()
page = writer.add_blank_page(width=200, height=300)
font = DictionaryObject({NameObject("/Type"): NameObject("/Font"),
                         NameObject("/Subtype"): NameObject("/Type1"), NameObject("/BaseFont"): NameObject("/Helvetica")})
page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})})
content = DecodedStreamObject()
content.set_data(b"BT /F1 12 Tf 20 200 Td (Hello document) Tj ET")
page[NameObject("/Contents")] = writer._add_object(content)
writer.write(text_pdf)
text_read = call("read", text_pdf, pages=[1])
assert "Hello document" in text_read["blocks"][0]["text"]
assert text_read["needs_ocr"] is False
call("read", "relative.docx", fail=True)
print(json.dumps({"word": str(word), "excel": str(excel), "pdf": str(pdf)}, ensure_ascii=False))
